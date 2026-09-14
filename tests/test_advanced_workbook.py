from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from advanced_tracking.workbook import (
    AdvancedWorkbookData,
    BATTED_BALL_COLUMNS,
    SCHEMA_VERSION,
    write_advanced_workbook,
)


class AdvancedWorkbookTests(unittest.TestCase):
    def test_writes_the_public_sheets_without_a_template(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "Game - Advanced.xlsx"
            write_advanced_workbook(
                target,
                AdvancedWorkbookData(
                    game={"session_id": "game-1", "away_team": "A", "home_team": "B"},
                ),
            )

            workbook = load_workbook(target, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                ["Game Summary", "Batting", "Pitching", "Fielding & Running",
                 "Batted Balls", "Fielding", "Baserunning", "Throws",
                 "Pitches", "Capture Health"],
            )
            self.assertEqual(workbook["Game Summary"]["B2"].value, SCHEMA_VERSION)
            workbook.close()

    def test_preserves_numeric_values_and_neutralizes_formula_text(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "Game - Advanced.xlsx"
            write_advanced_workbook(
                target,
                AdvancedWorkbookData(
                    game={"session_id": "game-2", "away_team": "=BAD", "home_team": "B"},
                    batted_balls=[{
                        "session_id": "game-2",
                        "play": 1,
                        "contact_timer": 1234,
                        "exit_velocity_mph": 101.25,
                    }],
                ),
            )

            workbook = load_workbook(target, data_only=False, read_only=True)
            game_values = list(workbook["Game Summary"].values)
            self.assertEqual(game_values[1][6], "'=BAD")
            ball_values = list(workbook["Batted Balls"].values)
            exit_index = BATTED_BALL_COLUMNS.index("exit_velocity_mph")
            self.assertEqual(ball_values[1][exit_index], 101.25)
            workbook.close()

    def test_replaces_an_existing_advanced_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "Game - Advanced.xlsx"
            write_advanced_workbook(target, AdvancedWorkbookData(game={"session_id": "old"}))
            write_advanced_workbook(target, AdvancedWorkbookData(game={"session_id": "new"}))

            workbook = load_workbook(target, read_only=True)
            self.assertEqual(workbook["Game Summary"]["A2"].value, "new")
            workbook.close()


class BattingSummaryTests(unittest.TestCase):
    def test_a_foul_is_not_a_ball_in_play(self):
        from advanced_tracking.summary import batting_summary

        rows = batting_summary([
            {"batter": "King Boo", "result": "fair_in_play",
             "measurement_status": "measured", "exit_velocity_mph": 100.0},
            {"batter": "King Boo", "result": "home_run",
             "measurement_status": "measured", "exit_velocity_mph": 106.0},
            {"batter": "King Boo", "result": "foul",
             "measurement_status": "measured", "exit_velocity_mph": 40.0},
            {"batter": "King Boo", "result": "foul_home_run",
             "measurement_status": "measured", "exit_velocity_mph": 42.0},
            {"batter": "King Boo", "result": "unknown",
             "measurement_status": "unavailable"},
        ])

        row = rows[0]
        self.assertEqual(row["balls_in_play"], 2)
        self.assertEqual(row["fouls"], 2)
        self.assertEqual(row["home_runs"], 1)
        self.assertEqual(row["batted_balls"], 5)
        self.assertEqual(row["measured_of_in_play"], "2/2")
        # The two fouls were weak; averaging them in would have said 72.0.
        self.assertEqual(row["avg_exit_velocity_mph"], 103.0)
        self.assertEqual(row["max_exit_velocity_mph"], 106.0)


if __name__ == "__main__":
    unittest.main()
