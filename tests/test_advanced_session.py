from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from advanced_tracking.session import AdvancedSession, advanced_workbook_path


class AdvancedSessionTests(unittest.TestCase):
    def test_finalizes_beside_original_and_removes_private_journal(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            original = output / "Fireballs vs Monsters.xlsx"
            session = AdvancedSession(output, session_id="session-1")
            self.assertTrue(session.record_play({
                "contact_timer": 20,
                "batter": "Mario",
                "batted_ball_class": "fair_in_play",
                "fielders": {},
                "runners": {},
                "throws": [],
            }, batted_ball={"exit_velocity_mph": 99.5}))
            self.assertFalse(session.record_play({"contact_timer": 20}))
            session.record_health({"capture_status": "complete", "frames_captured": 600})

            advanced = session.finalize(original, {"away_team": "Fireballs", "home_team": "Monsters"})
            self.assertEqual(advanced, advanced_workbook_path(original))
            self.assertTrue(advanced.exists())
            self.assertFalse(session.journal_path.exists())

            workbook = load_workbook(advanced, read_only=True)
            self.assertEqual(workbook["Batted Balls"]["A2"].value, 1)
            self.assertEqual(workbook["Batted Balls"]["J2"].value, 99.5)
            self.assertEqual(workbook["Batting"]["A2"].value, "Mario")
            self.assertEqual(workbook["Batting"]["B2"].value, 1)
            self.assertEqual(workbook["Batting"]["E2"].value, 99.5)
            workbook.close()

    def test_recovers_a_session_without_duplicating_records(self):
        with tempfile.TemporaryDirectory() as directory:
            first = AdvancedSession(directory, session_id="recover-me")
            first.record_play({
                "contact_timer": 42,
                "fielders": {},
                "runners": {},
                "throws": [],
            })
            recovered = AdvancedSession(directory, session_id="recover-me", recover=True)
            self.assertFalse(recovered.record_play({"contact_timer": 42}))
            result = recovered.finalize(Path(directory) / "Original.xlsx", {})

            workbook = load_workbook(result, read_only=True)
            self.assertEqual(workbook["Batted Balls"].max_row, 2)
            workbook.close()

    def test_pitch_event_does_not_claim_unmeasured_trajectory(self):
        with tempfile.TemporaryDirectory() as directory:
            session = AdvancedSession(directory, session_id="pitch-session")
            session.record_pitch({
                "pitch_timer": 10,
                "pitcher": "Mario",
                "batter": "Luigi",
                "offer": "take",
                "outcome": "strike",
            })
            advanced = session.finalize(Path(directory) / "Original.xlsx", {})
            workbook = load_workbook(advanced, read_only=True, data_only=True)
            headers = [cell.value for cell in workbook["Pitches"][1]]
            values = [cell.value for cell in workbook["Pitches"][2]]
            row = dict(zip(headers, values))
            self.assertEqual(row["measurement_status"], "event_only")
            self.assertEqual(row["invalid_reason"], "trajectory_metrics_unavailable")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
