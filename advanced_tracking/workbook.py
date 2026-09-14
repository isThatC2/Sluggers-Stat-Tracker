"""Create the additive advanced-stat workbook.

The community tracker's existing template and workbook writer deliberately do
not import this module. Advanced tracking produces a second workbook so a
failure here can never damage the original box score.
"""

from __future__ import annotations

import json
import os
import tempfile
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .summary import BATTING_COLUMNS, DEFENSE_COLUMNS, PITCHING_COLUMNS


SCHEMA_VERSION = "2.0"

# Columns whose numbers are only readable rounded. Excel stores the full value;
# this is the display format, so a sort or a formula still sees what was
# measured.
NUMBER_FORMATS = {
    1: ("route_efficiency", "avg_route_efficiency"),
    2: ("exit_velocity_mph", "max_exit_velocity_mph", "avg_exit_velocity_mph",
        "launch_angle_deg", "avg_launch_angle_deg", "spray_angle_deg",
        "distance_ft", "avg_distance_ft", "max_distance_ft", "path_ft",
        "displacement_ft", "lead_ft", "velocity_mph", "avg_velocity_mph",
        "max_velocity_mph", "peak_velocity_mph", "top_throw_mph",
        "sprint_speed_fps", "top_fielding_speed_fps", "top_running_speed_fps",
        "hang_time_s", "avg_hang_time_s", "flight_time_s"),
    3: ("reaction_s", "avg_reaction_s", "possession_to_release_s",
        "home_to_first_s", "best_home_to_first_s", "ninety_foot_split_s",
        "horizontal_break_units", "vertical_break_units",
        "avg_horizontal_break_units", "avg_vertical_break_units",
        "endpoint_x", "endpoint_y", "endpoint_z", "start_x", "start_y",
        "start_z", "touch_x", "touch_y", "touch_z", "end_x", "end_y", "end_z"),
}
DECIMALS = {name: f"0.{'0' * places}"
            for places, names in NUMBER_FORMATS.items() for name in names}

GAME_COLUMNS = (
    "session_id", "schema_version", "tracker_version", "memory_profile",
    "started_utc", "ended_utc", "away_team", "home_team", "park",
    "time_of_day", "innings", "ended_early",
)
# Detail columns are ordered so the first screen is the answer and the raw
# coordinates sit at the far right, where they are available without being in
# the way.
BATTED_BALL_COLUMNS = (
    "play", "inning", "inning_half", "outs", "balls", "strikes", "batter",
    "pitcher", "result", "exit_velocity_mph", "launch_angle_deg",
    "spray_angle_deg", "spray_side", "distance_ft", "hang_time_s",
    "endpoint_type", "measurement_status", "invalid_reason", "contact_timer",
    "endpoint_x", "endpoint_y", "endpoint_z",
)
FIELDING_COLUMNS = (
    "play", "fielder", "position", "is_primary", "first_touch", "fielded",
    "caught_in_flight", "catch_type", "path_ft", "displacement_ft",
    "route_efficiency", "sprint_speed_fps", "reaction_s", "route_valid",
    "route_excluded_because", "character_id", "start_x", "start_y", "start_z",
    "touch_x", "touch_y", "touch_z",
)
BASERUNNING_COLUMNS = (
    "play", "runner", "origin_base", "destination_base", "bases_ran",
    "lead_ft", "path_ft", "sprint_speed_fps", "home_to_first_s",
    "ninety_foot_split_s", "stealing", "movement_valid", "invalid_reason",
    "character_id",
)
THROW_COLUMNS = (
    "play", "sequence", "thrower", "thrower_position", "receiver",
    "receiver_position", "intended_target_position", "target_base",
    "peak_velocity_mph", "possession_to_release_s", "flight_time_s",
    "buddy_throw", "outs_recorded", "measurement_status", "invalid_reason",
    "start_x", "start_y", "start_z", "end_x", "end_y", "end_z",
)
PITCH_COLUMNS = (
    "pitch", "inning", "inning_half", "outs", "balls_before",
    "strikes_before", "batter", "pitcher", "pitch_in_pa", "velocity_mph",
    "horizontal_break_units", "vertical_break_units", "offer", "outcome",
    "measurement_status", "invalid_reason", "pitch_timer",
)
HEALTH_COLUMNS = (
    "session_id", "capture_status", "calibration_status", "frames_captured",
    "frames_missed", "missed_frame_rate", "position_offset",
    "fielder_pointers_valid", "advanced_export_status", "notes",
)


@dataclass(slots=True)
class AdvancedWorkbookData:
    """Normalized rows ready for an advanced workbook."""

    game: Mapping[str, Any]
    batted_balls: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    fielding: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    baserunning: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    throws: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    pitches: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    capture_health: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    batting: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    pitching: Sequence[Mapping[str, Any]] = field(default_factory=tuple)
    defense: Sequence[Mapping[str, Any]] = field(default_factory=tuple)


def _excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (list, tuple, dict)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    else:
        value = str(value)
    # Team names are user-editable. Do not let one become an Excel formula.
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _append_sheet(
    workbook: Workbook,
    title: str,
    columns: Sequence[str],
    rows: Iterable[Mapping[str, Any]],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"

    widths = [len(column) for column in columns]
    for row in rows:
        values = [_excel_value(row.get(column)) for column in columns]
        sheet.append(values)
        for index, value in enumerate(values):
            if value is not None:
                widths[index] = min(48, max(widths[index], len(str(value))))
    for index, (column, width) in enumerate(zip(columns, widths), start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width + 2
        decimals = DECIMALS.get(column)
        if decimals:
            for cell in sheet[letter][1:]:
                cell.number_format = decimals


def write_advanced_workbook(
    destination: Path | str,
    data: AdvancedWorkbookData,
) -> Path:
    """Write *destination* atomically and return its resolved path."""

    destination = Path(destination).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    game = dict(data.game)
    game.setdefault("schema_version", SCHEMA_VERSION)
    # Summaries first: they are what gets opened, and the per-play sheets
    # behind them are what they are made of.
    _append_sheet(workbook, "Game Summary", GAME_COLUMNS, (game,))
    _append_sheet(workbook, "Batting", BATTING_COLUMNS, data.batting)
    _append_sheet(workbook, "Pitching", PITCHING_COLUMNS, data.pitching)
    _append_sheet(workbook, "Fielding & Running", DEFENSE_COLUMNS, data.defense)
    _append_sheet(workbook, "Batted Balls", BATTED_BALL_COLUMNS, data.batted_balls)
    _append_sheet(workbook, "Fielding", FIELDING_COLUMNS, data.fielding)
    _append_sheet(workbook, "Baserunning", BASERUNNING_COLUMNS, data.baserunning)
    _append_sheet(workbook, "Throws", THROW_COLUMNS, data.throws)
    _append_sheet(workbook, "Pitches", PITCH_COLUMNS, data.pitches)
    _append_sheet(workbook, "Capture Health", HEALTH_COLUMNS, data.capture_health)

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=destination.parent,
            prefix=destination.stem + ".",
            suffix=".tmp.xlsx",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, destination)
    finally:
        workbook.close()
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()
    return destination
